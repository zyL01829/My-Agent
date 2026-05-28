from __future__ import annotations

import hashlib
import math
import shutil
from pathlib import Path
from typing import Any, Dict, List

import chromadb
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from app.database import DATA_DIR, from_json, get_conn, now_local, rows_to_dicts, to_json


UPLOAD_DIR = DATA_DIR / "uploads"
CHROMA_DIR = DATA_DIR / "chroma"


class HashEmbeddingFunction:
    def __init__(self, dimensions: int = 384) -> None:
        self.dimensions = dimensions

    def name(self) -> str:
        return "local_hash_embedding"

    def __call__(self, input: List[str]) -> List[List[float]]:
        return [self._embed(text) for text in input]

    def embed_query(self, input: List[str]) -> List[List[float]]:
        return self(input)

    def embed_documents(self, input: List[str]) -> List[List[float]]:
        return self(input)

    def _embed(self, text: str) -> List[float]:
        vector = [0.0] * self.dimensions
        tokens = [token for token in text.lower().replace("\n", " ").split(" ") if token]
        for token in tokens:
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            idx = int.from_bytes(digest[:4], "big") % self.dimensions
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[idx] += sign
        norm = math.sqrt(sum(v * v for v in vector)) or 1.0
        return [v / norm for v in vector]


class RAGService:
    def __init__(self) -> None:
        CHROMA_DIR.mkdir(parents=True, exist_ok=True)
        self.client = chromadb.PersistentClient(path=str(CHROMA_DIR))
        self.collection = self.client.get_or_create_collection(
            name="personal_knowledge",
            embedding_function=HashEmbeddingFunction(),
            metadata={"description": "Local personal steward knowledge base"},
        )

    def upload_file(self, source_path: Path, original_name: str, tags: List[str] | None = None) -> Dict[str, Any]:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        suffix = Path(original_name).suffix.lower()
        stored_name = f"{hashlib.sha1((original_name + str(source_path.stat().st_mtime)).encode()).hexdigest()[:12]}{suffix}"
        stored_path = UPLOAD_DIR / stored_name
        shutil.copyfile(source_path, stored_path)

        text_blocks = self._extract_text(stored_path, suffix)
        chunks = self._chunk_blocks(text_blocks)
        with get_conn() as conn:
            cur = conn.execute(
                """
                INSERT INTO knowledge_files(filename, stored_path, file_type, tags, chunk_count, uploaded_at)
                VALUES(?, ?, ?, ?, ?, ?)
                """,
                (original_name, str(stored_path), suffix.lstrip(".") or "unknown", to_json(tags or []), len(chunks), now_local()),
            )
            file_id = cur.lastrowid
            ids = []
            docs = []
            metas = []
            for index, chunk in enumerate(chunks):
                vector_id = f"file-{file_id}-chunk-{index}"
                metadata = {"file_id": file_id, "filename": original_name, "chunk_index": index, **chunk["metadata"]}
                conn.execute(
                    """
                    INSERT INTO knowledge_chunks(file_id, chunk_index, content, metadata, vector_id, created_at)
                    VALUES(?, ?, ?, ?, ?, ?)
                    """,
                    (file_id, index, chunk["content"], to_json(metadata), vector_id, now_local()),
                )
                ids.append(vector_id)
                docs.append(chunk["content"])
                metas.append(metadata)
            if ids:
                self.collection.upsert(ids=ids, documents=docs, metadatas=metas)

        return {"file_id": file_id, "filename": original_name, "chunk_count": len(chunks), "tags": tags or []}

    def list_files(self) -> List[Dict[str, Any]]:
        with get_conn() as conn:
            rows = rows_to_dicts(conn.execute("SELECT * FROM knowledge_files ORDER BY uploaded_at DESC").fetchall())
        for row in rows:
            row["tags"] = from_json(row["tags"], [])
        return rows

    def delete_file(self, file_id: int) -> Dict[str, Any]:
        with get_conn() as conn:
            file_row = conn.execute("SELECT * FROM knowledge_files WHERE id = ?", (file_id,)).fetchone()
            if not file_row:
                return {"ok": False, "message": "文件不存在。"}
            chunk_rows = conn.execute("SELECT vector_id FROM knowledge_chunks WHERE file_id = ?", (file_id,)).fetchall()
            vector_ids = [row["vector_id"] for row in chunk_rows]
            conn.execute("DELETE FROM knowledge_files WHERE id = ?", (file_id,))
        if vector_ids:
            self.collection.delete(ids=vector_ids)
        Path(file_row["stored_path"]).unlink(missing_ok=True)
        return {"ok": True, "deleted_file_id": file_id, "deleted_chunks": len(vector_ids)}

    def query(self, query: str, limit: int = 5) -> Dict[str, Any]:
        if not query.strip():
            return {"answer": "请输入要检索的问题。", "sources": []}
        result = self.collection.query(query_texts=[query], n_results=limit)
        documents = result.get("documents", [[]])[0]
        metadatas = result.get("metadatas", [[]])[0]
        distances = result.get("distances", [[]])[0] if result.get("distances") else [None] * len(documents)
        sources = []
        for doc, meta, distance in zip(documents, metadatas, distances):
            sources.append(
                {
                    "content": doc,
                    "filename": meta.get("filename", "unknown"),
                    "chunk_index": meta.get("chunk_index"),
                    "score": None if distance is None else round(float(distance), 4),
                }
            )
        answer = self._compose_answer(query, sources)
        return {"answer": answer, "sources": sources}

    def _compose_answer(self, query: str, sources: List[Dict[str, Any]]) -> str:
        if not sources:
            return "知识库中暂未检索到相关资料。"
        lines = [f"基于知识库中与“{query}”最相关的资料，建议优先参考："]
        for source in sources[:3]:
            snippet = source["content"].replace("\n", " ")[:180]
            lines.append(f"- {source['filename']}：{snippet}")
        return "\n".join(lines)

    def _extract_text(self, path: Path, suffix: str) -> List[Dict[str, Any]]:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return [
                {"content": page.extract_text() or "", "metadata": {"page": index + 1}}
                for index, page in enumerate(reader.pages)
            ]
        if suffix == ".docx":
            doc = Document(str(path))
            return [{"content": paragraph.text, "metadata": {"paragraph": index + 1}} for index, paragraph in enumerate(doc.paragraphs)]
        if suffix == ".pptx":
            prs = Presentation(str(path))
            blocks = []
            for slide_index, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
                blocks.append({"content": "\n".join(texts), "metadata": {"slide": slide_index + 1}})
            return blocks
        if suffix in {".md", ".txt"}:
            return [{"content": path.read_text(encoding="utf-8", errors="ignore"), "metadata": {"section": "document"}}]
        return [{"content": path.read_text(encoding="utf-8", errors="ignore"), "metadata": {"section": "document"}}]

    def _chunk_blocks(self, blocks: List[Dict[str, Any]], chunk_size: int = 900, overlap: int = 120) -> List[Dict[str, Any]]:
        chunks: List[Dict[str, Any]] = []
        for block in blocks:
            text = " ".join(block["content"].split())
            if not text:
                continue
            start = 0
            while start < len(text):
                chunk = text[start : start + chunk_size]
                chunks.append({"content": chunk, "metadata": block["metadata"]})
                if start + chunk_size >= len(text):
                    break
                start += chunk_size - overlap
        return chunks


rag_service = RAGService()
