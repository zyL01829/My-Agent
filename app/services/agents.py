from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Inches, Pt

from app.database import DATA_DIR, from_json, get_conn, now_local, rows_to_dicts, to_json
from app.services.model_router import choose_model
from app.services.notifications import notifications
from app.services.rag import rag_service


GENERATED_DIR = DATA_DIR / "generated"


AGENT_REGISTRY = [
    {
        "agent_id": "supervisor",
        "name": "主管 Agent",
        "description": "任务分发、上下文管理、模型路由和工具权限控制。",
        "enabled": True,
        "tools": ["model_router", "task_store", "notification_center", "rag"],
    },
    {
        "agent_id": "learning",
        "name": "学习规划 Agent",
        "description": "生成学习路线、每日计划、复盘问题，并根据进度调整计划。",
        "enabled": True,
        "tools": ["rag", "notification_center", "scheduler"],
    },
    {
        "agent_id": "ppt",
        "name": "PPT 制作 Agent",
        "description": "根据主题、风格和知识库素材生成 PPT 文件。",
        "enabled": True,
        "tools": ["rag", "pptx_generator", "file_store"],
    },
    {
        "agent_id": "frontier",
        "name": "前沿资讯 Agent",
        "description": "生成 AI 与金融量化领域的每日热点简报。",
        "enabled": True,
        "tools": ["web_search_adapter", "notification_center", "archive"],
    },
    {
        "agent_id": "paper",
        "name": "论文 Agent（预留）",
        "description": "后续实现论文逻辑、语言、图表和数据审查。",
        "enabled": False,
        "tools": ["rag", "document_parser"],
    },
    {
        "agent_id": "travel",
        "name": "旅游 Agent（预留）",
        "description": "后续实现旅游路线、实时天气、交通和机酒建议。",
        "enabled": False,
        "tools": ["web_search_adapter", "weather_adapter", "notification_center"],
    },
]


class SupervisorAgent:
    def route(self, message: str) -> Dict[str, Any]:
        lowered = message.lower()
        if "ppt" in lowered or "幻灯" in message:
            target = "ppt"
        elif "学习" in message or "计划" in message:
            target = "learning"
        elif "ai" in lowered or "量化" in message or "前沿" in message:
            target = "frontier"
        else:
            target = "learning"
        model = choose_model("supervisor", "cheap")
        return {"target_agent": target, "model": model.__dict__, "message": message}


class LearningAgent:
    def create_plan(
        self,
        topic: str,
        duration_days: int,
        daily_minutes: int,
        target_level: str,
        use_knowledge: bool = True,
    ) -> Dict[str, Any]:
        model = choose_model("learning", "default", uses_rag=use_knowledge)
        rag_context = rag_service.query(topic, limit=3) if use_knowledge else {"sources": []}
        phases = self._build_phases(topic, duration_days, target_level)
        days = []
        for day in range(1, duration_days + 1):
            phase = phases[min(len(phases) - 1, int((day - 1) / max(1, duration_days / len(phases))))]
            days.append(
                {
                    "day": day,
                    "title": f"{topic} Day {day}: {phase['name']}",
                    "goal": phase["goal"],
                    "minutes": daily_minutes,
                    "tasks": [
                        f"用 {max(15, daily_minutes // 3)} 分钟复习上一日关键点",
                        f"学习“{phase['name']}”核心概念并做结构化笔记",
                        "完成 3 个自测问题，记录不确定内容",
                    ],
                    "check": "用自己的话解释今日概念，并标记 0-100 的掌握度。",
                }
            )
        plan = {
            "topic": topic,
            "target_level": target_level,
            "duration_days": duration_days,
            "daily_minutes": daily_minutes,
            "phases": phases,
            "days": days,
            "rag_sources": rag_context.get("sources", []),
            "model": model.__dict__,
        }
        with get_conn() as conn:
            cur = conn.execute(
                """
                INSERT INTO learning_plans(topic, duration_days, daily_minutes, target_level, plan, created_at, updated_at)
                VALUES(?, ?, ?, ?, ?, ?, ?)
                """,
                (topic, duration_days, daily_minutes, target_level, to_json(plan), now_local(), now_local()),
            )
            plan_id = cur.lastrowid
            conn.execute(
                """
                INSERT INTO tasks(agent_id, title, status, input, output, model_used, created_at, updated_at)
                VALUES(?, ?, ?, ?, ?, ?, ?, ?)
                """,
                ("learning", f"学习计划：{topic}", "completed", to_json({"topic": topic}), to_json(plan), model.model, now_local(), now_local()),
            )
        notifications.create(
            "新的学习计划已生成",
            f"{topic} 的 {duration_days} 天计划已经准备好，今日任务会显示在首页。",
            "success",
            "learning",
            {"plan_id": plan_id},
        )
        return {"id": plan_id, **plan}

    def record_progress(self, plan_id: int, day_index: int, completion: int, mastery: int, notes: str = "") -> Dict[str, Any]:
        with get_conn() as conn:
            conn.execute(
                """
                INSERT INTO learning_progress(plan_id, day_index, completion, mastery, notes, created_at)
                VALUES(?, ?, ?, ?, ?, ?)
                """,
                (plan_id, day_index, completion, mastery, notes, now_local()),
            )
            row = conn.execute("SELECT * FROM learning_plans WHERE id = ?", (plan_id,)).fetchone()
            plan = from_json(row["plan"], {}) if row else {}
        adjustment = "按原计划继续。"
        if completion < 70 or mastery < 70:
            adjustment = "建议明天减少新内容，增加复习、错题整理和概念复述。"
        notifications.create("学习进度已记录", adjustment, "info", "learning", {"plan_id": plan_id, "day": day_index})
        return {"plan_id": plan_id, "day_index": day_index, "adjustment": adjustment, "plan": plan}

    def list_plans(self) -> List[Dict[str, Any]]:
        with get_conn() as conn:
            rows = rows_to_dicts(conn.execute("SELECT * FROM learning_plans ORDER BY created_at DESC").fetchall())
        for row in rows:
            row["plan"] = from_json(row["plan"], {})
        return rows

    def _build_phases(self, topic: str, duration_days: int, target_level: str) -> List[Dict[str, str]]:
        return [
            {"name": "建立知识地图", "goal": f"理解 {topic} 的核心概念、边界和常见应用。"},
            {"name": "系统学习与练习", "goal": f"围绕 {target_level} 目标完成主题学习、练习和案例分析。"},
            {"name": "项目化输出", "goal": f"用作品、总结或讲解验证自己是否达到 {target_level}。"},
        ]


class PPTAgent:
    def generate(self, topic: str, audience: str, style: str, slides: int, use_knowledge: bool = True) -> Dict[str, Any]:
        model = choose_model("ppt", "strong", needs_file=True, uses_rag=use_knowledge)
        rag_context = rag_service.query(topic, limit=5) if use_knowledge else {"sources": []}
        outline = self._outline(topic, audience, slides, rag_context.get("sources", []))
        GENERATED_DIR.mkdir(parents=True, exist_ok=True)
        filename = f"ppt_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        path = GENERATED_DIR / filename
        self._write_pptx(path, topic, audience, style, outline)
        output = {
            "topic": topic,
            "audience": audience,
            "style": style,
            "slides": outline,
            "file": f"/generated/{filename}",
            "model": model.__dict__,
            "sources": rag_context.get("sources", []),
        }
        with get_conn() as conn:
            conn.execute(
                """
                INSERT INTO tasks(agent_id, title, status, input, output, model_used, created_at, updated_at)
                VALUES(?, ?, ?, ?, ?, ?, ?, ?)
                """,
                ("ppt", f"PPT：{topic}", "completed", to_json({"topic": topic, "audience": audience}), to_json(output), model.model, now_local(), now_local()),
            )
        notifications.create("PPT 已生成", f"{topic} 的 PPT 已完成，可在 PPT 页面下载。", "success", "ppt", output)
        return output

    def list_jobs(self) -> List[Dict[str, Any]]:
        with get_conn() as conn:
            rows = rows_to_dicts(conn.execute("SELECT * FROM tasks WHERE agent_id = 'ppt' ORDER BY created_at DESC").fetchall())
        for row in rows:
            row["input"] = from_json(row["input"], {})
            row["output"] = from_json(row["output"], {})
        return rows

    def _outline(self, topic: str, audience: str, slides: int, sources: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        base = [
            ("主题与目标", f"面向 {audience} 说明 {topic} 的背景、价值和本次汇报目标。"),
            ("核心框架", f"用 3-5 个模块建立 {topic} 的整体结构。"),
            ("关键内容", "结合知识库资料提炼概念、方法、案例或数据。"),
            ("实践路径", "给出可执行步骤、资源需求和风险点。"),
            ("总结与下一步", "凝练结论，并明确后续行动。"),
        ]
        outline = []
        for index in range(slides):
            title, body = base[index] if index < len(base) else (f"补充分析 {index + 1}", f"继续展开 {topic} 的细节。")
            citation = sources[index % len(sources)]["filename"] if sources else "暂无知识库引用"
            outline.append({"slide": index + 1, "title": title, "body": body, "citation": citation})
        return outline

    def _write_pptx(self, path: Path, topic: str, audience: str, style: str, outline: List[Dict[str, Any]]) -> None:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = topic
        title_slide.placeholders[1].text = f"{audience} | {style}"
        for item in outline:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = item["title"]
            body = slide.placeholders[1].text_frame
            body.clear()
            p = body.paragraphs[0]
            p.text = item["body"]
            p.font.size = Pt(22)
            cite = body.add_paragraph()
            cite.text = f"引用来源：{item['citation']}"
            cite.font.size = Pt(12)
        prs.save(path)


class FrontierAgent:
    def run_daily(self) -> Dict[str, Any]:
        model = choose_model("frontier", "default")
        today = date.today().isoformat()
        report = {
            "date": today,
            "sections": [
                {
                    "name": "AI 模型与产品",
                    "items": [
                        "关注主流模型更新、开源模型发布、Agent 工具链变化。",
                        "建议后续接入新闻搜索 API 后替换为实时来源。",
                    ],
                },
                {
                    "name": "金融与量化",
                    "items": [
                        "关注宏观数据、市场波动、量化平台工具和监管动态。",
                        "本地原型阶段生成结构化待查清单，避免伪造实时新闻。",
                    ],
                },
                {
                    "name": "行动建议",
                    "items": ["将高价值来源沉淀到知识库，周末做趋势复盘。"],
                },
            ],
            "model": model.__dict__,
            "sources": ["web_search_adapter:reserved"],
        }
        with get_conn() as conn:
            existing = conn.execute("SELECT id FROM frontier_reports WHERE report_date = ?", (today,)).fetchone()
            if existing:
                conn.execute("UPDATE frontier_reports SET content = ? WHERE id = ?", (to_json(report), existing["id"]))
                report_id = existing["id"]
            else:
                cur = conn.execute(
                    "INSERT INTO frontier_reports(report_date, title, content, created_at) VALUES(?, ?, ?, ?)",
                    (today, f"{today} 前沿日报", to_json(report), now_local()),
                )
                report_id = cur.lastrowid
        notifications.create("今日前沿日报已生成", "AI 与金融量化日报已更新到前沿页面。", "info", "frontier", {"report_id": report_id})
        return {"id": report_id, **report}

    def list_reports(self) -> List[Dict[str, Any]]:
        with get_conn() as conn:
            rows = rows_to_dicts(conn.execute("SELECT * FROM frontier_reports ORDER BY report_date DESC LIMIT 30").fetchall())
        for row in rows:
            row["content"] = from_json(row["content"], {})
        return rows


supervisor_agent = SupervisorAgent()
learning_agent = LearningAgent()
ppt_agent = PPTAgent()
frontier_agent = FrontierAgent()
